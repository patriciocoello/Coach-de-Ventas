"""PLANTILLA: licitación (Ej. 27 / Ej. 28) en el formato SENCILLO del usuario.

    python ejemplo_licitacion.py <carpeta_salida>

Una fila = un experimento: Mi costo | Mi oferta (RiskSimtable) | Oferta competidor 1..n |
Mi Ganancia = RiskOutput()+IF(MIN(ofertas)=mi oferta, mi oferta - costo, 0).
Debajo: tabla "Posibles Ofertas | Ganancia Promedio", fila ganadora en amarillo + "IDEAL".
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from montecarlo import np, simtable, mejor  # noqa: E402
from risk_excel import LibroRisk, DINERO  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

salida_dir = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
N_EXP = 500
COSTO = 100_000
N_COMP = 4
A, B = 100_000, 300_000                        # RiskUniform(A;B) de cada competidor
OFERTAS = list(range(110_000, 250_001, 10_000))

# ---------------------------------------------------------------- réplica Python
def modelo(oferta, n, rng):
    comp = rng.uniform(A, B, (n, N_COMP))
    return np.where(oferta <= comp.min(axis=1), oferta - COSTO, 0.0)

r_n = simtable(OFERTAS, modelo, N_EXP)          # lo que se parecerá a @RISK
r_big = simtable(OFERTAS, modelo, 20_000)       # para confirmar el óptimo
print("Óptimo (500 exp.):", mejor(r_n), " Óptimo (20,000):", mejor(r_big))

# ---------------------------------------------------------------- Excel
lb = LibroRisk("Ej. 28 Estás por entrar a un concurso ... (enunciado completo)",
               iteraciones=N_EXP, simulaciones=len(OFERTAS), rango_enunciado="B2:L10")
fila = 15
lb.v(f"B{fila - 1}", "a)")
encab = ["Experimento", "Mi costo", "Mi oferta "] + [f"Oferta competidor {i + 1}" for i in range(N_COMP)] + ["Mi Ganancia"]
for j, h in enumerate(encab):
    lb.v(f"{get_column_letter(2 + j)}{fila}", h)
d = fila + 1
col_ult_comp = get_column_letter(4 + N_COMP)
col_gan = get_column_letter(5 + N_COMP)
tabla_ini, tabla_fin = d + 3, d + 3 + len(OFERTAS) - 1
lb.v(f"B{d}", 1)
lb.v(f"C{d}", COSTO, formato=DINERO)
lb.f(f"D{d}", f"=RiskSimtable(C{tabla_ini}:C{tabla_fin})", formato=DINERO)
for i in range(N_COMP):
    col = get_column_letter(5 + i)
    lb.f(f"{col}{d}", f"=RiskUniform({A},{B})", formato=DINERO, celda_texto=f"{col}{fila - 1}")
lb.f(f"{col_gan}{d}", f"=RiskOutput()+IF(MIN(D{d}:{col_ult_comp}{d})=D{d},D{d}-C{d},0)",
     formato=DINERO, celda_texto=f"{col_gan}{fila - 1}")
lb.tabla_simtable(f"C{d + 2}", OFERTAS, "Posibles Ofertas", "Ganancia Promedio",
                  resultados=[round(v) for _, v in r_n], paso=10_000)
lb.conclusion(f"F{d + 6}", f"Lo ideal es ofrecer {mejor(r_n)[0]:,}, ahi la ganancia promedio es la mas alta",
              ancho_cols=4, alto_filas=2)
lb.anchos(16, {"A": 5, "B": 14})
print("Excel:", lb.guardar(str(salida_dir / "Ejercicio28_plantilla.xlsx")))
