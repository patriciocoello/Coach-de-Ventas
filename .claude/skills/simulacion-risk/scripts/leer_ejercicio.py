"""Lee un ejercicio en Excel (como los de clase) y muestra:
- el enunciado (está en cuadros de texto, no en celdas),
- las celdas con valores y fórmulas (las de @RISK incluidas),
- las imágenes incrustadas (se extraen a una carpeta para poder verlas),
- los ajustes de simulación guardados (iteraciones, simulaciones).

Uso:  python -I leer_ejercicio.py ejercicio.xlsx [carpeta_imagenes]
"""
import re
import sys
import zipfile
from pathlib import Path

import openpyxl


def enunciados(ruta):
    textos = []
    with zipfile.ZipFile(ruta) as z:
        for nombre in sorted(n for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n)):
            xml = z.read(nombre).decode("utf8", "ignore")
            for bloque in re.split(r"</xdr:(?:twoCellAnchor|oneCellAnchor|absoluteAnchor)>", xml):
                parrafos = re.findall(r"<a:p>(.*?)</a:p>", bloque, re.S)
                t = "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", p)) for p in parrafos).strip()
                pos = re.search(r"<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>", bloque, re.S)
                donde = f"col {int(pos.group(1)) + 1}, fila {int(pos.group(2)) + 1}" if pos else "?"
                if t:
                    t = t.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&")
                    textos.append((donde, t))
                elif "<xdr:pic>" in bloque:
                    textos.append((donde, "[IMAGEN]"))
    return textos


def extraer_imagenes(ruta, destino):
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    salidas = []
    with zipfile.ZipFile(ruta) as z:
        for n in z.namelist():
            if n.startswith("xl/media/"):
                p = destino / Path(n).name
                p.write_bytes(z.read(n))
                salidas.append(str(p))
    return salidas


def celdas(ruta, max_celdas=1500):
    wb = openpyxl.load_workbook(ruta)
    wv = openpyxl.load_workbook(ruta, data_only=True)
    for hoja in wb.sheetnames:
        if hoja.startswith(("RiskSerialization", "rsklib")):
            continue  # datos internos de @RISK
        ws, vs = wb[hoja], wv[hoja]
        print(f"\n=== Hoja {hoja} ({ws.dimensions})")
        vistos, n = set(), 0
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is None:
                    continue
                val = c.value.text if hasattr(c.value, "text") else c.value
                # colapsa filas repetidas (tablas de 500 experimentos)
                patron = (c.column_letter, re.sub(r"[A-Z]*\d+", "#", str(val))[:150])
                if isinstance(val, str) and val.startswith("=") and patron in vistos:
                    continue
                vistos.add(patron)
                s = repr(val)
                s = s if len(s) < 300 else s[:300] + "..."
                calc = vs[c.coordinate].value
                print(f"{c.coordinate}: {s}" + (f"  -> {calc!r}" if isinstance(val, str) and val.startswith("=") else ""))
                n += 1
                if n >= max_celdas:
                    print("... (truncado)")
                    return
    ajustes = {k: wb.defined_names[k].attr_text for k in ("RiskNumIterations", "RiskNumSimulations")
               if k in wb.defined_names}
    if ajustes:
        print("\nAjustes @RISK:", ajustes)


if __name__ == "__main__":
    ruta = sys.argv[1]
    print("### ENUNCIADO / CUADROS DE TEXTO")
    for donde, t in enunciados(ruta):
        print(f"[{donde}] {t}\n")
    imgs = extraer_imagenes(ruta, sys.argv[2] if len(sys.argv) > 2 else Path(ruta).stem + "_imagenes")
    if imgs:
        print("### IMÁGENES extraídas (ábrelas: suelen ser la matriz/tabla del enunciado o el histograma de @RISK):")
        print("\n".join(imgs))
    print("\n### CELDAS")
    celdas(ruta)
